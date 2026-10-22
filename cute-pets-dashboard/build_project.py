import json,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
SP='/tmp/claude-0/-home-user-onetake/8d7ced99-b727-5fab-ae82-daa300c2d5fb/scratchpad'
OUT='/home/user/onetake/ملف-مشروع-كلود/بيانات-كيوت-بيتس-2026.xlsx'
d=json.load(open(SP+'/proj.json'))
N={'exhibit':'المعرض','cpstore':'كيوت بيتس','hamtaro':'هامتارو','wagy':'واجي','wholesale':'الجملة','main':'الرئيسي'}
C=['exhibit','cpstore','hamtaro','wagy','wholesale']
MN={1:'يناير',2:'فبراير',3:'مارس',4:'أبريل',5:'مايو',6:'يونيو',7:'يوليو',8:'أغسطس',9:'سبتمبر',10:'أكتوبر',11:'نوفمبر',12:'ديسمبر'}
ml=lambda m: MN[int(m[5:])]+' '+m[:4]
r2=lambda v: round(v,2) if isinstance(v,(int,float)) else v
lab=lambda a: re.sub(r'\s+',' ',re.sub(r'\s*\([^)]*\)','',str(a))).strip()
H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='0F6B5C')
wb=Workbook(); first=[True]
def sheet(title,head,rows):
    ws=wb.active if first[0] else wb.create_sheet(); first[0]=False
    ws.title=title; ws.sheet_view.rightToLeft=True; ws.append(head)
    for c in ws[1]: c.font=H; c.fill=F; c.alignment=Alignment(horizontal='center')
    for r in rows: ws.append([r2(v) for v in r])
    for i,h in enumerate(head): ws.column_dimensions[ws.cell(1,i+1).column_letter].width=max(12,min(44,len(str(h))+6))
    ws.freeze_panes='A2'; return ws
ms=sorted(d)
# summary
ws=sheet('ملخص المتاجر',['المركز']+[ml(m) for m in ms]+['إجمالي السنة'],[])
ws.delete_rows(1)
def block(title,key,sign=0):
    ws.append([title]); ws.cell(ws.max_row,1).font=Font(bold=True,size=13)
    ws.append(['المركز']+[ml(m) for m in ms]+['إجمالي السنة'])
    for c in ws[ws.max_row]: c.font=H; c.fill=F
    tots=[0]*len(ms)
    for cid in C:
        vals=[round(key(d[m][cid]['x']),2) for m in ms]; tots=[a+b for a,b in zip(tots,vals)]; ws.append([N[cid]]+vals+[round(sum(vals),2)])
    if sign:
        ex=[round(sign*d[m]['main']['x'].get('tax',0),2) for m in ms]; ws.append(['فرق ضريبة الشركة']+ex+[round(sum(ex),2)]); tots=[a+b for a,b in zip(tots,ex)]
    ws.append(['إجمالي الشركة']+[round(t,2) for t in tots]+[round(sum(tots),2)])
    for c in ws[ws.max_row]: c.font=Font(bold=True)
    ws.append([])
block('الإيرادات',lambda x:x.get('revenue',0)); block('صافي الربح قبل الضريبة',lambda x:x.get('net',0))
block('الضريبة',lambda x:x.get('tax',0),1); block('صافي الربح بعد الضريبة',lambda x:x.get('netAfter',x.get('net',0)),-1)
ws.column_dimensions['A'].width=20
for row in ws.iter_rows():
    for x in row:
        if isinstance(x.value,(int,float)): x.number_format='#,##0;[Red]-#,##0'
# goals (latest month)
lm=ms[-1]; D=d[lm]
HM=re.compile('رمل|مكاف|مستلزمات|العاب|ألعاب|خداش|بيوت|اكواخ|اقفاص')
SG={'exhibit':(200000,8000),'cpstore':(170000,17000),'hamtaro':(280000,35000),'wagy':(75000,7500)}
MG={'cpstore':(500,320,34,22),'hamtaro':(920,285,38,25),'wagy':(260,270,29,22)}
rows=[]
for c in ['exhibit','cpstore','hamtaro','wagy']:
    x=D[c]['x']; ads=D[c]['summary']['ads']; rev=x['revenue']; st,at=SG[c]; n=N[c]
    rows+=[[n,'الفرع','المبيعات (ريال)','لا يقل عن',st,rev],[n,'الفرع','صرف الإعلان (ريال)','لا يزيد عن',at,ads],[n,'الفرع','كل ريال إعلان يجيب (ريال مبيعات)','لا يقل عن',st/at,rev/ads if ads else 0]]
    if c in MG:
        o,aov,cpo,mx=MG[c]; od=x.get('ordAll',0); cats=D[c]['cats']; t=sum(max(0,v) for v in cats.values()); hm=sum(max(0,v) for k,v in cats.items() if HM.search(k))
        rows+=[[n,'التسويق','عدد الطلبات','لا يقل عن',o,od],[n,'التسويق','متوسط قيمة الطلب بدون التوصيل (ريال)','لا يقل عن',aov,(rev-x.get('delivery',0))/od if od else 0],[n,'التسويق','تكلفة الإعلان لكل طلب (ريال)','لا يزيد عن',cpo,ads/od if od else 0],[n,'التسويق','مبيعات الرمل والمكافآت والمستلزمات (% من المبيعات)','لا يقل عن',mx,hm/t*100 if t else 0]]
sheet('أهداف الشهر',['الفرع','الجهة','الهدف','الشرط','المستهدف','الفعلي '+ml(lm),'الحالة'],[r+['محقق' if (r[5]>=r[4] if r[3]=='لا يقل عن' else r[5]<=r[4]) else 'غير محقق'] for r in rows])
# monthly results
rows=[]
for m in ms:
    for c in C:
        x=d[m][c]['x']; s=d[m][c]['summary'] or {}
        rows.append([ml(m),m,N[c],x.get('revMain',0),x.get('revBrand',0),x.get('delivery',0),x.get('otherRev',0),x.get('funRev',0),x.get('groomRev',0),x.get('revenue',0),x.get('cogs',0),x.get('gross',0),x.get('adminT',0),x.get('totalExp',0)-x.get('adminT',0)-x.get('cogs',0),x.get('netBeforeMgr',x.get('net',0)),x.get('mgrComm',0),x.get('net',0),x.get('tax',0),x.get('netAfter',x.get('net',0)),s.get('variable',0),s.get('fixed',0),s.get('ads',0)])
    t=d[m]['main']['x'].get('tax',0); rows.append([ml(m),m,'فرق ضريبة الشركة (على الرئيسي)']+[0]*12+[t,-t,0,0,0])
sheet('النتائج الشهرية',['الشهر','الرمز','المركز','مبيعات بدون الجملة','مبيعات براندات الجملة','إيراد التوصيل','إيرادات أخرى','منها الترفيه (منطقة المرح)','منها الجرومنج','إجمالي الإيرادات','تكلفة البضاعة','مجمل الربح','المصاريف الإدارية','المصاريف التشغيلية (بدون التكلفة)','صافي قبل عمولة المدير','عمولة المدير','صافي قبل الضريبة','الضريبة','صافي بعد الضريبة','تكاليف متغيرة','تكاليف ثابتة','الدعاية'],rows)
rows=[]
for m in ms:
    rev=sum(d[m][c]['x']['revenue'] for c in C); g=sum(d[m][c]['x']['gross'] for c in C); n=sum(d[m][c]['x']['net'] for c in C)
    a=sum(d[m][c]['x'].get('netAfter',d[m][c]['x']['net']) for c in C)-d[m]['main']['x'].get('tax',0); rows.append([ml(m),rev,g,n,n-a,a])
sheet('إجمالي الشركة',['الشهر','الإيرادات','مجمل الربح','صافي قبل الضريبة','الضريبة','صافي بعد الضريبة'],rows)
rows=[]; CAT={'admin':'إدارية','expense':'تشغيلية','other':'إيرادات أخرى'}
for m in ms:
    for c in C:
        x=d[m][c]['x']
        for k,l in [('إيرادات المبيعات بدون الجملة','revMain'),('إيرادات مبيعات الجملة','revBrand'),('إيراد التوصيل داخل الرياض','delIn'),('خدمات الشحن الخارجي','delOut'),('تكلفة البضاعة المباعة','cogsMain'),('تكلفة بضاعة الجملة','cogsWs'),('محروقات (موزعة بالطلبات)','fuelAlloc'),('صيانة سيارات (موزعة بالطلبات)','vehAlloc'),('مواد تغليف (موزعة بالطلبات)','packAlloc')]:
            if abs(x.get(l,0))>0.004: rows.append([ml(m),m,N[c],'إيراد' if l in('revMain','revBrand','delIn','delOut') else 'تشغيلية',k,x[l],''])
        for l in d[m][c]['lines']:
            if l.get('cat') in ('sales','contra','delivery') or abs(l['amount'])<0.004: continue
            rows.append([ml(m),m,N[c],CAT.get(l['cat'],l['cat']),lab(l['account']),l['amount'],'تقديري' if l.get('estimate') else ''])
sheet('بنود قائمة الدخل',['الشهر','الرمز','المركز','النوع','البند','المبلغ','ملاحظة'],rows)
rows=[]
for m in ms:
    for c in C:
        tot=sum(d[m][c]['cats'].values()) or 1
        for k,v in sorted(d[m][c]['cats'].items(),key=lambda kv:-kv[1]):
            cost=d[m][c]['catCost'].get(k); rows.append([ml(m),m,N[c],k.strip(),v,round(v/tot*100,2),cost if cost is not None else '',(v-cost) if cost is not None else '',round((v-cost)/v*100,2) if cost is not None and v else ''])
sheet('مبيعات الأقسام',['الشهر','الرمز','المركز','القسم','المبيعات','النسبة من مبيعات المركز %','التكلفة','مجمل الربح','هامش مجمل الربح %'],rows)
# break-even (as on the dashboard): variable = cost of goods + selling costs that move with sales; fixed = the rest (ads included)
rows=[]
for m in ms:
    for c in C:
        if c=='wholesale': continue
        b=d[m][c]['summary']; rev=b['revenue']; var=b['variable']; fx=b['fixed']; ads=b['ads']
        if not rev: continue
        cm=rev-var; cmr=cm/rev; be=fx/cmr if cmr>0 else None
        rows.append([ml(m),m,N[c],rev,var,round(var/rev*100,2),cm,round(cmr*100,2),fx,ads,be if be else 'لا يتعادل',round(rev/be*100,1) if be else '',(rev-be) if be else '',b['beNet'],round(ads/rev*100,2),round(rev/ads,2) if ads else ''])
sheet('نقطة التعادل ومؤشرات الإعلان',['الشهر','الرمز','المركز','الإيرادات','التكاليف المتغيرة','المتغيرة % من الإيرادات','هامش المساهمة','نسبة هامش المساهمة %','التكاليف الثابتة','منها الدعاية','إيرادات نقطة التعادل','الإيرادات ÷ نقطة التعادل %','هامش الأمان (الإيرادات − التعادل)','صافي قبل العمولة والضريبة','الإعلان % من الإيرادات','كل ريال إعلان يجيب (MER)'],rows)
B={'kaniva':'كانيفا','ronron':'رون رون','sherman':'شيرمان','candy':'كاندي'}; rows=[]
for m in ms:
    for c,o in (d[m]['wholesale']['matrix'] or {}).items():
        for b,v in o.items(): rows.append([ml(m),m,N.get(c,c),B.get(b,b),v.get('qty',0),v.get('sales',0),v.get('wh',0),v.get('imp',0)])
sheet('البراندات حسب المركز',['الشهر','الرمز','المركز','البراند','الكمية','مبيعات المتجر للعميل','قيمة الجملة (سعر القائمة)','تكلفة الاستيراد'],rows)

# ad spend by cost center (Edara ledger of the ads account)
import os, collections
AL=SP+'/ads_ledger.json'
if os.path.exists(AL):
    A=json.load(open(AL)); MNl=lambda m: ml(m)
    sheet('الإعلانات - تفصيل القيود',['التاريخ','الشهر','المركز','المنصة','مركز التكلفة في إدارة','المبلغ','رقم القيد','البيان'],[[a['date'],MNl(a['month']),a['center'],a['platform'],a['cc'],a['amount'],a['doc'],a['note']] for a in A])
    agg=collections.defaultdict(float)
    for a in A: agg[(a['month'],a['center'],a['platform'])]+=a['amount']
    sheet('الإعلانات حسب المنصة',['الشهر','الرمز','المركز','المنصة','المبلغ'],[[MNl(k[0]),k[0],k[1],k[2],round(v,2)] for k,v in sorted(agg.items())])
    months=sorted({a['month'] for a in A}); plats=['جوجل','ميتا','سناب شات','تيك توك','واتساب','رواتب السوشيال ميديا','تحسين محركات البحث (SEO)','مباشر على المركز (بدون منصة)','كاش باك درب']
    ws=sheet('ملخص الإعلانات',['المركز','المنصة']+[MNl(m) for m in months]+['الإجمالي'],[])
    for cen in ['هامتارو','كيوت بيتس','واجي','المعرض','الجملة','الرئيسي']:
        for p_ in plats:
            vals=[round(agg.get((m,cen,p_),0),2) for m in months]
            if any(vals): ws.append([cen,p_]+vals+[round(sum(vals),2)])
        tot=[round(sum(v for k,v in agg.items() if k[0]==m and k[1]==cen),2) for m in months]
        if any(tot):
            ws.append([cen,'إجمالي المركز']+tot+[round(sum(tot),2)])
            for c in ws[ws.max_row]: c.font=Font(bold=True)
    tot=[round(sum(v for k,v in agg.items() if k[0]==m),2) for m in months]; ws.append(['الشركة','إجمالي الإعلانات']+tot+[round(sum(tot),2)])
    for c in ws[ws.max_row]: c.font=Font(bold=True)
    for row in ws.iter_rows():
        for x in row:
            if isinstance(x.value,(int,float)): x.number_format='#,##0;[Red]-#,##0'

# online orders by city and by Riyadh district (from Salla's orders): order count and delivery revenue
GM=SP+'/orders26/geo_months.json'; GS=SP+'/geo/orders_geo_sep.json'
if os.path.exists(GM):
    G=json.load(open(GM))
    if os.path.exists(GS): G['2026-09']=json.load(open(GS))
    for key,title,col in [('cities','الطلبات حسب المدينة','المدينة'),('districts','الطلبات حسب الحي (الرياض)','الحي')]:
        rows=[]
        for m in sorted(G):
            for e in sorted(G[m].get(key,[]),key=lambda e:-sum(e['n'].values())):
                for c,n in e['n'].items():
                    if n: rows.append([ml(m),m,N[c],e['name'],n,round(e['f'].get(c,0),2)])
        sheet(title,['الشهر','الرمز','المركز',col,'عدد الطلبات','إيراد التوصيل'],rows)
# the dashboard's other report tables, as shown on the page for every month (cards.json from cards.mjs)
CJ=SP+'/cards.json'
if os.path.exists(CJ):
    import re
    D=json.load(open(CJ)); last=max(D)
    def num(v):
        t=v.replace(',','').replace('−','-').strip()
        return float(t) if re.fullmatch(r'-?\d+(\.\d+)?',t) else v
    def isHead(r): return not any(isinstance(num(v),float) for v in r[1:])
    MONTHNAME=re.compile(r'^(يناير|فبراير|مارس|أبريل|مايو|يونيو|يوليو|أغسطس|سبتمبر|أكتوبر|نوفمبر|ديسمبر|الإجمالي \()')
    STORE={'كيوت':'cpstore','هامتارو':'hamtaro','واجي':'wagy','المحل':'exhibit','المعرض':'exhibit','محمد عاطف':'hamtaro'}
    def who(cell):
        for k,c in STORE.items():
            if k in cell: return N[c]
        return 'كل المتاجر'
    # company-page tables: the center is named in the row
    for title,sh in [('العمولات المستحقة عن الشهر','العمولات المستحقة'),('التوصيل ونسب التحمل','التوصيل ونسب التحمل')]:
        rows=[]; head=None
        for m in sorted(D):
            t=D[m]['all'].get(title)
            if not t: continue
            for r in t['rows']:
                if head is None and isHead(r): head=r; continue
                if isHead(r) or r[0]=='الإجمالي': continue
                rows.append([ml(m),m,who(' '.join(r[:2]))]+[num(v) for v in r])
        if head: sheet(sh,['الشهر','الرمز','المركز']+head,rows)
    # center-page tables: one sheet each, raw rows (a row without numbers is a sub-header)
    for title,sh in [('التوصيل والمصاريف المحمّلة','التوصيل والمصاريف المحمّلة'),('نقطة التعادل وربحية التصنيفات','التعادل وربحية التصنيفات'),
                     ('إيرادات الخدمات','إيرادات الخدمات - تفصيل'),('تقرير مبيعات الجملة','تقرير مبيعات الجملة'),
                     ('تكلفة بضاعة الجملة حسب القسم','تكلفة الجملة حسب القسم'),('أصناف بسعر جملة تقديري','أصناف بسعر جملة تقديري'),
                     ('مقارنة المراكز في بيع البراندات','مقارنة المراكز في البراندات')]:
        rows=[]; w=0
        for m in sorted(D):
            for c in C:
                t=D[m].get(c,{}).get(title)
                if not t: continue
                for k in t['kpis']: rows.append([ml(m),m,N[c],'مؤشر',k[0],num(k[1] or ''),k[2] or ''])
                for r in t['rows']:
                    if MONTHNAME.match(r[0]) and m!=last: continue   # the month-by-month history is kept once, from the last month
                    rows.append([ml(m),m,N[c],'عنوان' if isHead(r) else 'بيان']+[num(v) for v in r]); w=max(w,len(r))
        if rows: sheet(sh,['الشهر','الرمز','المركز','نوع الصف']+[f'عمود {i+1}' for i in range(max(w,3))],rows)
wb.save(OUT); print('saved',[w.title for w in wb.worksheets])
