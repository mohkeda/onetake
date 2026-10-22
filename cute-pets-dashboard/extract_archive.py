import openpyxl, warnings, json, sys, re
warnings.filterwarnings('ignore')
wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
STORES = {'المحل': 'exhibit', 'كيوت': 'cpstore', 'هامتارو': 'hamtaro', 'واجي': 'wagy'}
def sheet_of(store, n): return f'{store} ' if n == 1 else f'{store}  ({n})'
def num(v): return float(v) if isinstance(v, (int, float)) else 0.0
out = {}
for store, cid in STORES.items():
    for n in range(1, 9):
        ws = wb[sheet_of(store, n)]
        sec = None; rev = []; adm = []; op = []; net_cell = ws['D1'].value
        for r in range(2, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if not isinstance(a, str): continue
            t = a.strip()
            if t.startswith('أولا'): sec = 'rev'; continue
            if t.startswith('ثانيا'): sec = None; continue
            if t.startswith('المصاريف الادار'): sec = 'adm'; continue
            if t.startswith('المصاريف التشغي'): sec = 'op'; continue
            if t in ('الإجمالي', 'البيان', 'الاجمالي'): 
                if t != 'البيان': sec = None
                continue
            v = num(ws.cell(r, 2).value)
            if sec == 'rev': rev.append([t, round(v, 2)])
            elif sec == 'adm': adm.append([t, round(v, 2)])
            elif sec == 'op': op.append([t, round(v, 2)])
        tr, ta, to = sum(x[1] for x in rev), sum(x[1] for x in adm), sum(x[1] for x in op)
        net = tr - ta - to
        m = f'2026-{n:02d}'
        out.setdefault(m, {})[cid] = {'rev': rev, 'admin': adm, 'op': op}
        flag = '' if isinstance(net_cell, (int, float)) and abs(net - net_cell) < 1 else f'  !! sheet D1={net_cell}'
        print(f'{m} {cid:8} rev={tr:10.0f} adm={ta:9.0f} op={to:10.0f} net={net:9.0f}{flag}  title={ws["A1"].value}')
    # categories by month from the latest sheet
    ws = wb[sheet_of(store, 8)]
    hdr = None
    for row in ws.iter_rows(max_row=10):
        for c in row:
            if c.value == 'Jan': hdr = (c.row, c.column); break
        if hdr: break
    r0, c0 = hdr
    months = []
    for k in range(8):
        months.append(f'2026-{k+1:02d}')
    for r in range(r0 + 1, r0 + 40):
        lab = ws.cell(r, c0 - 1).value
        if not isinstance(lab, str) or lab.strip() in ('الإجمالي', 'الاجمالي'): break
        for k, m in enumerate(months):
            v = num(ws.cell(r, c0 + k).value)
            if v: out[m][cid].setdefault('cats', {})[lab.strip()] = round(v, 2)
json.dump(out, open(sys.argv[2], 'w'), ensure_ascii=False)
